"""AEM Sites: “On this page” must list the headings down to H4, and each page is checked for breaking issues
and picture problems (not loaded, stretched, outside its area, sideways scroll)."""
import pymupdf

from pdfval import load_config
from pdfval.engine import compare_url
from test_site_nav import make_site, serve  # noqa: F401  (the fixture)


def _png(path, w, h):
    pm = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pm.clear_with(180)
    pm.save(str(path))


def _site(tmp_path, base, extra_html):
    pdf = make_site(tmp_path)
    page = tmp_path / "guide" / "overview.html"
    _png(tmp_path / "guide" / "pic.png", 200, 100)
    page.write_text(page.read_text().replace('<div class="pager">', extra_html + '<div class="pager">'))
    cfg = load_config()
    cfg["site"]["check_otp"] = True  # otp is off by default (old check) - this file tests it directly
    return compare_url(str(pdf), base + "overview.html", str(tmp_path / "run"), cfg)["site"]


def rows(site, group, status="fail"):
    return [r for r in site["rows"] if r["group"] == group and r["status"] == status]


def test_h4_missing_from_on_this_page_is_reported(tmp_path, serve):
    site = _site(tmp_path, serve, '<h3 id="x1">Panel</h3><p>About the panel.</p><h4 id="x2">Touch layer</h4><p>About the touch layer.</p>')
    missing = {r["item"]: r["note"] for r in rows(site, "otp")}
    assert missing.get("Touch layer") == "h4 not listed" and missing.get("Panel") == "h3 not listed"


def test_picture_and_layout_breaks_are_reported(tmp_path, serve):
    site = _site(tmp_path, serve,
                 '<p>Good picture:</p><img src="pic.png" alt="good" width="200" height="100">'
                 '<p>Stretched:</p><img src="pic.png" alt="stretched" style="width:200px;height:200px">'
                 '<p>Too wide:</p><img src="pic.png" alt="wide" style="width:900px;height:450px;max-width:none">'
                 '<p>Missing:</p><img src="nope.png" alt="gone" width="100" height="50">')
    got = {(r["item"], r["actual"]) for r in rows(site, "layout")}
    assert ("Picture stretched", "stretched") in got
    assert ("Picture outside its area", "wide") in got
    assert ("Picture not loaded", "gone") in got
    assert not any(a == "good" for _, a in got)


def test_clean_page_has_no_layout_failure(tmp_path, serve):
    # centred (design spec: every picture is centred) - an un-styled inline <img> would otherwise sit left,
    # which is exactly what check_image_center now correctly flags; this test is about a genuinely clean page
    site = _site(tmp_path, serve, '<p>Picture:</p><img src="pic.png" alt="ok" width="200" height="100" style="display:block;margin:0 auto">')
    assert not rows(site, "layout") and rows(site, "layout", "pass")


def test_left_aligned_picture_fails_the_design_spec(tmp_path, serve):
    """Design spec ("Image rules": Alignment = Center): a picture left in its column is reported even
    though it is the only picture on the page (check_picture_alignment's relative/majority check would
    see nothing to disagree with - this is the separate, absolute check_image_center)."""
    site = _site(tmp_path, serve, '<p>Picture:</p><img src="pic.png" alt="left" width="200" height="100">')
    bad = rows(site, "layout")
    assert any(r["item"] == "Picture alignment (design spec)" and r["expected"] == "centred" and "left" in r["actual"] for r in bad)


_STEP_TEXT = ("1. Connect the USB cable between the PC and the monitor via the upstream USB port. This "
             "upstream USB port transmits data between the PC and the USB devices connected to the monitor.")


def test_numbered_step_without_hanging_indent_is_reported(tmp_path, serve):
    """A step number typed as plain text ("1. ...", not a real <li> counter) that wraps to a second line
    must hang-indent: the wrapped line has to start to the right of the "1.", clearing it - not flush
    beneath it, which prints the next line's text directly under the number instead of under the step's
    own text."""
    site = _site(tmp_path, serve, f'<p>{_STEP_TEXT}</p>')
    bad = rows(site, "layout")
    assert any(r["item"] == "Wrapped line not indented under the text" for r in bad), bad


def test_numbered_step_with_hanging_indent_is_not_reported(tmp_path, serve):
    """The same wrapped step, this time with the standard hanging-indent CSS trick (padding-left +
    negative text-indent) applied: the wrapped line clears the marker, so nothing is reported."""
    site = _site(tmp_path, serve, f'<p style="padding-left:1.5em;text-indent:-1.5em">{_STEP_TEXT}</p>')
    assert not any(r["item"] == "Wrapped line not indented under the text" for r in rows(site, "layout"))


_STICKY = '<div style="position:fixed;top:0;left:0;right:0;height:60px;background:#fff;z-index:10">Sticky bar</div>'
# enough height that scrolling the heading below it to the top of the viewport is actually possible
# (the page must be taller than the viewport by more than the bar's height, or the browser clamps the
# scroll short of fully aligning the heading, and the bar never gets the chance to cover it)
_FILLER = '<div style="height:2000px">spacer</div>'


def test_heading_under_sticky_header_is_reported(tmp_path, serve):
    """A heading far down the page (so jumping to it, like a reader clicking it in the nav, actually
    scrolls): if a sticky/fixed top bar stays pinned over the viewport and the page does not reserve
    space for it, the jump lands with the heading's top partly covered by the bar - reported even though
    nothing is wrong at the page's own load scroll position, which this bug never shows at."""
    # a spacer after the heading too: the browser can only scroll a heading all the way to the top of the
    # viewport if there is enough page left below it to make that scroll position reachable at all
    site = _site(tmp_path, serve, _STICKY + _FILLER + '<h2 id="jump">Remote controller exterior view</h2>' + _FILLER)
    bad = rows(site, "layout")
    hit = next((r for r in bad if r["item"] == "Jump target hidden under sticky header"
                and r["actual"] == "Remote controller exterior view"), None)
    assert hit is not None, bad


def test_heading_clear_of_sticky_header_is_not_reported(tmp_path, serve):
    """The same sticky bar, but the heading reserves space for it (scroll-margin-top matching the bar's
    height, the standard fix): the browser's own scrollIntoView honours it, so the jump lands clear of
    the bar - not reported."""
    site = _site(tmp_path, serve, _STICKY + _FILLER + '<h2 id="jump" style="scroll-margin-top:64px">Clear heading</h2>' + _FILLER)
    assert not any(r["item"] == "Jump target hidden under sticky header" and r["actual"] == "Clear heading"
                   for r in rows(site, "layout"))


def test_empty_on_this_page_box_does_not_swallow_the_page(tmp_path, serve):
    """A topic with no sub-headings renders its "On this page" box empty - no links of its own. The panel
    search must not climb out of that empty box and up past the content root looking for one, or the
    page's whole container is taken for the panel and every heading and paragraph in it counts as chrome:
    the typography checks then sample nothing at all and report a clean page instead of its real issues."""
    site = _site(tmp_path, serve, '<div class="toc topic-toc-root is-toc-empty"><p>On this page</p></div>')
    assert not [r for r in site["rows"] if "nothing found to sample" in r.get("actual", "")]
    # the content really was sampled: the page's own H1 was measured against the design spec
    assert [r for r in site["rows"] if r["group"] == "typography" and r["item"].startswith("Headline 1")]



def test_a_design_spec_issue_is_reported_with_the_picture(tmp_path, serve):
    """A picture left in its column fails the design spec. The bug report must show the picture, not only
    say "centred -> left": the reader has to see the spot to judge and to fix it."""
    from docx import Document

    from pdfval.report import writer
    pdf = make_site(tmp_path)
    page = tmp_path / "guide" / "overview.html"
    _png(tmp_path / "guide" / "pic.png", 200, 100)
    page.write_text(page.read_text().replace(
        '<div class="pager">', '<p>Picture:</p><img src="pic.png" alt="left" width="200" height="100"><div class="pager">'))
    run_dir = tmp_path / "run"
    result = compare_url(str(pdf), serve + "overview.html", str(run_dir), load_config())
    rows = [r for r in result["site"]["rows"] if r["item"] == "Picture alignment (design spec)"]
    assert rows and rows[0].get("shot"), rows
    writer.write_pdf_report(result, run_dir)
    doc = Document(run_dir / "genuine-issues.docx")
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Bug")
    assert table.rows[0].cells[5].text == "On the page"
    blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    assert any(c.paragraphs[0].runs and c.paragraphs[0].runs[0]._element.findall(f".//{blip}")
               for r in table.rows[1:] for c in [r.cells[5]]), "no design-spec row carries its picture"
