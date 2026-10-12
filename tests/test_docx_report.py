"""The Word report holds the issues of the PDF report - the same ones - and nothing else: no metrics, no tables."""
import re

import pytest

docx = pytest.importorskip("docx")

from pdfval.report import docx_report, pdf_report


def _finding(fid, issue, message):
    return {"id": fid, "check": "content", "category": "content", "severity": "warning", "message": message, "description": message,
            "issue": issue, "genuine": True, "types": ["missing text"], "baseline": [{"page": 3, "bbox": [0, 0, 10, 10]}], "candidate": [],
            "detail": {}, "critical": False}


def test_same_issues_as_the_pdf_report_and_no_tables(tmp_path):
    section = {"id": 0, "title": "Getting started", "status": "fail", "findings": [
        _finding("000-001", "Data missing", "Missing text: “Press OK”"),
        _finding("000-002", "Extra content", "Extra text: “Press Cancel”"),
        {**_finding("000-003", "", "a note"), "genuine": False}]}
    result = {"meta": {"name": "Test product"}, "sections": [section], "summary": {}}
    expected = [f["id"] for _, f in pdf_report.select_issues(result, pdf_report.GENUINE["filter"], None)]
    doc = docx.Document(str(docx_report.build(result, tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert re.findall(r"#(\d{3}-\d{3})", text) == expected == ["000-001", "000-002"]
    assert not doc.tables
    assert "Getting started" in text and "Score" not in text


def test_word_gap_shows_the_space_count_on_both_sides(tmp_path):
    """"Timing  Support" (2 spaces) -> "Timing Support" (1 space): once the spaces are collapsed - and
    _q collapses them, as Word does with a run of spaces - the two sides read exactly alike, so the
    count has to be in the text. Without it the report printed the same line twice and the reader
    could not see what was wrong."""
    from pdfval.genuine import prod_stage
    f = {"check": "content", "category": "content", "issue": "Space after a word differs",
         "types": ["spacing"], "severity": "warning", "genuine": True,
         "message": "Word gap differs: “Timing  Support” (2 spaces) → “Timing Support” (1 space)",
         "detail": {"op": "spacing", "baseline_spaces": 2, "candidate_spaces": 1, "words": 1}}
    prod, stage = prod_stage(f)
    assert prod != stage, (prod, stage)
    assert "2 spaces" in prod and "1 space" in stage


def test_issue_without_a_prod_stage_pair_still_says_what_differs(tmp_path):
    """A font weight off the design spec, a table column whose rows start at different edges: there is no
    "prod text → stage text" pair to print, and the .docx used to show nothing but the issue's name and
    its two screenshots. It now prints the same lines the PDF report does."""
    from pdfval.report.docx_report import _message_lines
    from pdfval.genuine import prod_stage
    f = {"check": "style", "category": "css", "issue": "Font weight off the design spec",
         "types": ["spec font-weight"], "severity": "warning", "genuine": True,
         "message": "Caption — font weight: Figma Regular (400), stage Bold (700)\n"
                    "Figma: Roboto Regular, 12pt / 16pt line height, colour #757575\n"
                    "Stage: Roboto Bold, 12pt, colour #000000",
         "description": "Caption — font weight: Figma Regular (400), stage Bold (700)\n"
                        "Figma: Roboto Regular, 12pt / 16pt line height, colour #757575\n"
                        "Stage: Roboto Bold, 12pt, colour #000000"}
    assert prod_stage(f) is None   # nothing to show as prod text vs stage text
    lines = _message_lines(f)
    assert [l[0] for l in lines] == ["Issue", "Figma", "Stage"], lines
    assert "Regular" in lines[1][1] and "Bold" in lines[2][1]


def test_every_difference_of_a_merged_finding_is_listed(tmp_path):
    """A renumbered list reports several differences at one spot ("13." → "3.", then the numbers prod has
    and stage does not). prod_stage describes only the first, so the .docx used to print that one and drop
    the rest, while the PDF report and the side-by-side viewer listed all of them."""
    import pymupdf
    from docx import Document
    from pdfval.report import docx_report
    f = {"check": "content", "category": "content", "issue": "Text changed + Bullet / list number missing",
         "types": ["changed text", "bullet marker"], "severity": "warning", "genuine": True, "id": "006-003",
         "message": "Changed text: “13.” → “3.”", "baseline": [], "candidate": [],
         "description": ("Changed text (2 places): “13.” → “3.”; “14.” → “4.”"
                         "  ·  List number missing in stage: “13.” before “ASPECT” in prod, none in stage"
                         "  ·  List number added in stage: “3.” before “Selects the display aspect”"),
         "detail": {"op": "replace", "baseline_text": "13.", "candidate_text": "3.", "words": 1,
                    "baseline_sentence": "13. ASPECT Selects the display aspect ratio.",
                    "candidate_sentence": "3. ASPECT Selects the display aspect ratio."}}
    result = {"meta": {}, "sections": [{"title": "Remote control", "findings": [f]}]}
    doc = Document(docx_report.build(result, tmp_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Difference 2" in text and "Difference 3" in text, text
    assert "ASPECT" in text and "Selects the display aspect" in text
