"""content.docx - the content review on its own.

It reports only: text missing from or added to stage, bold / italic on one side only, hyperlinks that
do not work or that arrive somewhere else, and how the text is aligned. Nothing about pictures, tables,
the TOC, fonts, colours, sizes, spacing or wrapping - see pdf_report.CONTENT_TYPES.

Every issue states Expected (prod) and Actual (stage), and the manuals' own text goes through verbatim,
whatever script it is written in.
"""
import pytest

from pdfval import engine
from pdfval.report import docx_report, pdf_report, writer

pytest.importorskip("docx")


def _finding(fid, types, message, expected="", actual="", **kw):
    return {"id": fid, "bug": f"Bug_{fid}", "check": kw.get("check", "content"),
            "category": kw.get("category", "content"), "severity": "warning", "message": message,
            "types": list(types), "genuine": True, "critical": False, "issue": kw.get("issue", types[0]),
            "baseline": [{"page": 0, "bbox": [0, 0, 10, 10]}], "candidate": [{"page": 0, "bbox": [0, 0, 10, 10]}],
            "detail": {"expected": expected, "actual": actual}, "color": "#333333"}


def _result(*findings):
    return {"meta": {"name": "Test product", "baseline": {"path": "/prod/x.pdf"}, "candidate": {"path": "/stage/x.pdf"}},
            "summary": {"content": {"match_pct": 99.0, "pass_pct": 95.0}},
            "sections": [{"id": "s1", "title": "Package contents", "findings": list(findings)}]}


def _text(path):
    from docx import Document
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def _heads(path):
    from docx import Document
    return [p.text for p in Document(str(path)).paragraphs if p.style.name.startswith(("Title", "Heading 1"))]


# --- what it reports ------------------------------------------------------------------------------

def test_the_types_it_covers_are_exactly_the_ones_asked_for():
    assert pdf_report.CONTENT_TYPES == {
        "missing text", "extra text", "missing section", "extra section",
        "emphasis",
        "broken link", "missing link", "link target differs", "link lands elsewhere",
        "spec text-align", "list alignment", "row alignment",
    }


def test_every_type_it_covers_is_actually_reported():
    """A type that is ignored or not genuine would be filtered out before it ever reached the report."""
    cfg = engine.load_config()
    genuine, ignored = set(cfg["genuine"]["types"]), set(cfg["ignore"]["types"])
    for t in pdf_report.CONTENT_TYPES:
        assert t not in ignored, f"{t} is in [ignore] types: never produced"
        assert t in genuine, f"{t} is not a genuine issue: the content report filters it out"


@pytest.mark.parametrize("types,bucket", [
    (["missing text"], 1), (["extra text"], 1), (["missing section"], 1),
    (["broken link"], 2), (["missing link"], 2), (["link target differs"], 2), (["link lands elsewhere"], 2),
    (["spec text-align"], 3), (["list alignment"], 3), (["row alignment"], 3),
    (["emphasis"], 4),
    (["image smaller"], 5),
])
def test_each_issue_lands_in_the_right_part(types, bucket):
    assert docx_report.content_bucket({"types": types}) == bucket


def test_an_issue_with_two_types_is_listed_once():
    assert docx_report.content_bucket({"types": ["emphasis", "missing text"]}) == 1


# --- the file -------------------------------------------------------------------------------------

def test_it_is_written_with_its_own_parts(tmp_path):
    out = writer.write_content_report(_result(
        _finding("001", ["missing text"], "Missing text", "“Press OK”", "Not in stage"),
        _finding("002", ["broken link"], "Broken link", "https://benq.com/a", "404", category="links"),
        _finding("003", ["emphasis"], "Bold differs", "plain in prod", "bold in stage", check="style", category="css"),
        _finding("004", ["list alignment"], "List numbers not aligned", "one left edge", "item 3 further in"),
    ), tmp_path)
    assert out and out.name == "content.docx"
    assert _heads(out) == ["Content report - Test product", "Content — missing and extra (1)",
                           "Hyperlinks — broken and redirected (1)", "Alignment (1)", "Bold / italic (1)"]


def test_it_leaves_out_everything_that_is_not_content(tmp_path):
    out = writer.write_content_report(_result(
        _finding("001", ["missing text"], "Missing text", "“Press OK”", "Not in stage"),
        _finding("002", ["image smaller"], "Image smaller in stage", category="images", check="assets"),
        _finding("003", ["missing row"], "Table row missing", category="tables", check="tables"),
        _finding("004", ["spec font-size"], "Font size differs", category="css", check="style"),
    ), tmp_path)
    text = _text(out)
    assert "missing text" in text and "Press OK" in text
    for gone in ("image smaller", "missing row", "spec font-size"):
        assert gone not in text, gone


def test_a_run_with_no_content_issue_writes_no_file(tmp_path):
    (tmp_path / "content.docx").write_bytes(b"stale")      # one left from an earlier run
    out = writer.write_content_report(_result(
        _finding("001", ["image smaller"], "Image smaller in stage", category="images", check="assets")), tmp_path)
    assert out is None and not (tmp_path / "content.docx").exists()


def test_every_issue_states_expected_and_actual(tmp_path):
    out = writer.write_content_report(_result(
        _finding("001", ["missing text"], "Missing text", "“Press the OK key”", "Not in stage")), tmp_path)
    text = _text(out)
    assert "Expected (prod):" in text and "Actual (stage):" in text
    assert "Press the OK key" in text


# --- all languages --------------------------------------------------------------------------------

@pytest.mark.parametrize("script,prod,stage", [
    ("Chinese", "按下 OK 鍵以確認設定", "按下 OK 鍵"),
    ("Japanese", "電源ボタンを押してください", "電源ボタンを押す"),
    ("Korean", "전원 버튼을 누르십시오", "전원 버튼"),
    ("Arabic", "اضغط على زر التشغيل للمتابعة", "اضغط على زر التشغيل"),
    ("Thai", "กดปุ่มเปิดปิดเพื่อดำเนินการต่อ", "กดปุ่มเปิดปิด"),
    ("Russian", "Нажмите кнопку питания", "Нажмите кнопку"),
    ("German", "Drücken Sie die Ein-/Aus-Taste größer", "Drücken Sie die Taste"),
])
def test_the_manuals_own_text_survives_in_any_script(tmp_path, script, prod, stage):
    """The report quotes the manual back; a guide in any language must come through unmangled."""
    out = writer.write_content_report(_result(
        _finding("001", ["missing text"], f"Missing text: {prod}", prod, stage)), tmp_path)
    text = _text(out)
    assert prod in text, (script, text[:300])
    assert stage in text, (script, text[:300])
