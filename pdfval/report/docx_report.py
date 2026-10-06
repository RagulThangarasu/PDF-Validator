"""The Word report (genuine-issues.docx): the issues of the PDF report (genuine-issues.pdf), one after the
other - the same issues in the same order, each with its section, its AEM topic and the prod / stage
screenshots with the highlights on both. The TOC is a table: prod and stage entries by level (L1-L3) with
the status of each entry (Pass / Fail / Extra on stage). Issues only: no metrics, no overview tables.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from . import pdf_report

SEV = {"error": "D92D20", "warning": "B45309", "info": "475569"}
# TOC status -> (label, text colour, cell fill)
TOC_STATUS = {"pass": ("Pass", "166534", "DCFCE7"), "fail": ("Fail", "B42318", "FEE2E2"), "extra": ("Extra on stage", "B45309", "FEF3C7")}
INK, MUTED, RULE = "1D2330", "6A7282", "CBD2DC"
# the report runs in priority order, 1 first; the picture issues come last
PRIORITIES = {1: "Content", 2: "Links and spacing", 3: "TOC, sections, tables and layout", 4: "Fonts and colours",
              5: "Images (size, pixelated, icons, extra images)"}
# the issues named by the reviewer, by label (f["issue"]); every other issue takes the priority of its group
ISSUE_PRIORITY = {"Text changed": 1, "Data missing + Extra content": 1, "Space after a word differs": 2,
                  "Table header not centred": 3, "Extra image": 5, "Image differs": 5}


def _priority(f: dict) -> int:
    """The priority of an issue in the report (see PRIORITIES)."""
    label = f.get("issue") or ", ".join(f.get("types", [])) or f["check"]
    if label in ISSUE_PRIORITY:
        return ISSUE_PRIORITY[label]
    if f.get("category") == "images" or set(f.get("types") or []) & pdf_report.PDF_IMAGE_TYPES:
        return 5
    return {"content": 1, "links": 2, "css": 4}.get(pdf_report.group_of(f), 3)


def _rgb(hex_: str):
    from docx.shared import RGBColor
    h = hex_.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _run(par, text: str, *, bold: bool = False, colour: str = INK, size: float = 9, italic: bool = False, highlight=None):
    from docx.shared import Pt
    r = par.add_run(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text))  # control characters are not valid in a .docx
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = _rgb(colour)
    if highlight is not None:
        r.font.highlight_color = highlight
    return r


def _par(doc, *, before: float = 0, after: float = 2, keep: bool = True):
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p


def _link(par, text: str, url: str):
    """A clickable link (python-docx has no helper for it)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0F766E"), ("w:u", "single"), ("w:sz", "16")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rpr.append(el)
    t = OxmlElement("w:t")
    t.text = text
    r.append(rpr)
    r.append(t)
    link.append(r)
    par._p.append(link)


def _shade(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell(cell, text: str = "", *, bold: bool = False, colour: str = INK, size: float = 8.5, fill: str | None = None,
          align_centre: bool = False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = 0
    if align_centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        _run(p, text, bold=bold, colour=colour, size=size)
    if fill:
        _shade(cell, fill)


def _header_row(row, *, repeat: bool = False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trpr = row._tr.get_or_add_trPr()
    if repeat:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)


def _toc_status(row: dict) -> tuple[str, str]:
    """(TOC status key, reason) for one row of the TOC comparison: Pass, Fail or Extra on stage."""
    if row["status"] == "extra in stage":
        return "extra", ""
    flags = [fl for fl in row.get("flags", []) if fl]
    names = {"order differs": "sequence wrong"}  # the entry is in a different place in the sequence
    reasons = ([names.get(row["status"], row["status"])] if row["status"] != "match" else []) + flags
    return ("fail", ", ".join(reasons)) if reasons else ("pass", "")


def _cover(doc, result: dict) -> None:
    """The cover page check: the sheet's title and version against what the cover prints (prod and stage)."""
    from .. import cover as cover_check
    try:
        c = cover_check.check(result)
    except Exception:  # no sheet / unreadable sheet: the cover is not checked, the rest of the report is still made
        c = None
    if c is None:
        return
    doc.add_heading("Cover page", level=1)
    if not c["fields"]:
        _run(_par(doc, after=4), "No sheet row matches this manual: the cover is not checked.", colour=MUTED)
        return
    p = _par(doc, after=4)
    _run(p, f"Sheet row {c['row']} · {c['model']}  ·  ", colour=MUTED, size=8)
    fails = sum(f["status"] == "fail" for f in c["fields"])
    _run(p, f"{fails} fail" if fails else "all pass", bold=True, colour=TOC_STATUS["fail" if fails else "pass"][1], size=8)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Field", "Expected (sheet)", "Prod cover", "Stage cover", "Status"]):
        _cell(cell, text, bold=True, colour="FFFFFF", fill="475569", align_centre=True)
    _header_row(table.rows[0], repeat=True)
    for f in c["fields"]:
        cells = table.add_row().cells
        _header_row(table.rows[-1])
        _cell(cells[0], f["field"], bold=True)
        _cell(cells[1], f["expected"])
        _cell(cells[2], "✓ printed" if f["prod"] == "ok" else f["prod"])
        _cell(cells[3], "✓ printed" if f["stage"] == "ok" else f["stage"], colour=TOC_STATUS["pass" if f["stage"] == "ok" else "fail"][1])
        label, colour, fill = TOC_STATUS["pass" if f["status"] == "pass" else "fail"]
        _cell(cells[4], label, bold=True, colour=colour, fill=fill)
    _par(doc, after=2)


def _toc(doc, result: dict, seq: int) -> int:
    """The TOC as a table of three columns: the prod entry, the stage entry and the status of each entry."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches

    t = result.get("toc") or {}
    if not t.get("rows"):
        return 0
    heading = t.get("heading") or {}
    head_b, head_c = heading.get("baseline") or "—", heading.get("candidate") or "—"
    head_differs = head_b.strip().lower() != head_c.strip().lower()
    statuses = [_toc_status(r) for r in t["rows"]]
    n = {k: sum(1 for s, _ in statuses if s == k) for k in TOC_STATUS}

    doc.add_heading(f"Table of contents ({n['fail']} fail · {n['extra']} extra on stage · {n['pass']} pass)", level=1)
    if head_differs:  # a numbered issue like the others: OPEN, expected (prod) and actual (stage), highlighted
        from docx.enum.text import WD_COLOR_INDEX
        p = _par(doc, before=6, after=0)
        _run(p, f"{seq + 1}.  ", bold=True, colour=INK, size=9)
        _run(p, "TOC", bold=True, colour="9333EA", size=8)
        _run(p, "  ·  ", colour=MUTED, size=8)
        _run(p, "heading differs", bold=True, colour="9333EA", size=9)
        _open_tag(p)
        left, right = _diff_words(head_b, head_c)
        for label, pieces, col in (("Expected (prod): ", left, "2563EB"), ("Actual (stage): ", right, "0F766E")):
            p = _par(doc, after=0)
            _run(p, label, bold=True, colour=col)
            _run(p, "“", colour=INK)
            for k, (text, changed) in enumerate(pieces):
                _run(p, (" " if k else "") + text, bold=changed, colour="B42318" if changed else INK)
            _run(p, "”", colour=INK)
        seq += 1

    # three columns only: the prod entry, the stage entry (each indented by its level) and the status
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(4.0), Inches(4.0), Inches(2.0)]
    _header_row(table.rows[0], repeat=True)
    for cell, (label, fill) in zip(table.rows[0].cells, (("Prod", "2563EB"), ("Stage", "0F766E"), ("Status", "475569"))):
        _cell(cell, label, bold=True, colour="FFFFFF", fill=fill, align_centre=True)

    def entry(e) -> str:
        if not e:
            return "—"
        lvl = min(max(int(e.get("level") or 1), 1), 3) - 1
        page = e.get("page")
        return "\u00a0" * 6 * lvl + f"{e['title']}  (p.{page if page is not None else '—'})"

    for row, (key, reason) in zip(t["rows"], statuses):
        cells = table.add_row().cells
        _header_row(table.rows[-1])
        _cell(cells[0], entry(row.get("baseline")), bold=False)
        _cell(cells[1], entry(row.get("candidate")), bold=False)
        label, colour, fill = TOC_STATUS[key]
        _cell(cells[2], f"{label} · {reason}" if reason else label, bold=True, colour=colour, fill=fill)

    for r in table.rows:
        for c, w in zip(r.cells, widths):
            c.width = w
    return seq


def _screenshots(doc, out: Path, f: dict, shots: dict, img_w):
    from docx.shared import Inches
    imgs = {side: pdf_report._jpeg_bytes(out / shots[side]) if shots.get(side) else None for side in ("baseline", "candidate")}
    if not any(imgs.values()):
        _par(doc, after=4, keep=False)
        return
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    cap = _par(doc, before=2, after=0)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cap, "   |   ".join(f"{lab} p.{shots.get(side + '_page', '?')}" for side, lab in (("baseline", "PROD"), ("candidate", "STAGE")) if imgs[side]),
         bold=True, colour=MUTED, size=7)
    # prod and stage side by side, centred, both at the same height so the two pictures line up
    sides = [side for side in ("baseline", "candidate") if imgs[side]]
    h = min([Inches(4.6)] + [img_w * imgs[side][1] for side in sides])
    p = _par(doc, after=6, keep=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, side in enumerate(sides):
        if k:
            p.add_run("      ")
        p.add_run().add_picture(io.BytesIO(imgs[side][0]), height=int(h))


def build(result: dict, out_dir: str | Path, filename: str = "genuine-issues.docx") -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    out = Path(out_dir)
    flt = pdf_report.GENUINE["filter"]
    issues = pdf_report.select_issues(result, flt, None)
    toc_in_own_part = bool((result.get("toc") or {}).get("rows")) and any(f["check"] == "toc" for _, f in issues)
    if toc_in_own_part:  # as in the PDF report: the TOC is compared entry by entry in its own part
        issues = [x for x in issues if x[1]["check"] != "toc"]

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = max(sec.page_width, sec.page_height), min(sec.page_width, sec.page_height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.5))
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)
    for style in ("Title", "Heading 1", "Heading 2", "Heading 3"):  # one typeface and ink colour for every heading
        doc.styles[style].font.name = "Arial"
        doc.styles[style].font.color.rgb = _rgb(INK)
    img_w = (sec.page_width - sec.left_margin - sec.right_margin - Inches(0.4)) / 2

    name = (result.get("meta") or {}).get("name") or ""
    title = doc.add_heading(f"PDF report - issues{' - ' + name if name else ''}", level=0)
    title.paragraph_format.space_after = Pt(4)

    buckets = {p: [] for p in PRIORITIES}
    for s, f in issues:
        buckets[_priority(f)].append((s, f))
    _cover(doc, result)
    seq = 0  # the issues are numbered in the order of the report
    n_toc = 0
    if toc_in_own_part:  # the TOC comes first, as a table
        seq = _toc(doc, result, seq)
        n_toc = 1
    for p, label in PRIORITIES.items():
        items = buckets[p]
        if not items:
            continue
        doc.add_heading(f"Priority {p} · {label}" + (f" ({len(items)})" if items else ""), level=1)
        current = None
        for s, f in items:
            if s is not current:
                current = s
                doc.add_heading(s["title"], level=2)
            seq += 1
            _issue(doc, out, f, img_w, seq)
    if not issues and not n_toc:
        _run(_par(doc), "No issues.", colour="16A34A")
    path = out / filename
    doc.save(str(path))
    return path


def _diff_words(expected: str, actual: str) -> tuple[list[tuple[str, bool]], list[tuple[str, bool]]]:
    """The expected (prod) and actual (stage) text as (text, changed) pieces: the words only prod has (missing in
    stage) and the words only stage has (extra in stage) are changed; the rest is the same in both."""
    import difflib
    ta, tb = expected.split(), actual.split()
    left, right = [], []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes():
        if op == "equal":
            left.append((" ".join(ta[i1:i2]), False))
            right.append((" ".join(tb[j1:j2]), False))
        else:
            if i2 > i1:
                left.append((" ".join(ta[i1:i2]), True))
            if j2 > j1:
                right.append((" ".join(tb[j1:j2]), True))
    return left, right


def _open_tag(par) -> None:
    """The OPEN status, at the end of the issue line: red and larger."""
    _run(par, "  -  ", colour=MUTED, size=9)
    _run(par, "(OPEN)", bold=True, colour="B42318", size=11)


def _expected_actual(f: dict) -> tuple[str, str] | None:
    """What prod has (expected) and what stage has (actual) for the issue, or None when it cannot be put that way."""
    from ..genuine import prod_stage
    ps = prod_stage(f)
    return (re.sub(r"\.{4,}", " ", ps[0]), re.sub(r"\.{4,}", " ", ps[1])) if ps else None


def _issue(doc, out: Path, f: dict, img_w, seq: int):
    from docx.enum.text import WD_COLOR_INDEX
    pa = f"p.{f['baseline'][0]['page'] + 1}" if f.get("baseline") else "—"
    pc = f"p.{f['candidate'][0]['page'] + 1}" if f.get("candidate") else "—"
    p = _par(doc, before=6, after=0)
    _run(p, f"{seq}.  ", bold=True, colour=INK, size=9)
    if f.get("critical"):
        _run(p, "CRITICAL  ·  ", bold=True, colour="B42318", size=8)
    _run(p, pdf_report.SEV_LABEL.get(f["severity"], f["severity"].upper()), bold=True, colour=SEV.get(f["severity"], "475569"), size=8)
    _run(p, "  ·  ", colour=MUTED, size=8)
    _run(p, f.get("issue") or ", ".join(f.get("types", [])) or f["check"], bold=True, colour=(f.get("color") or "#333333").lstrip("#"), size=9)
    _run(p, f"  ·  #{f['id']}  ·  prod {pa} ↔ stage {pc}", colour=MUTED, size=8)
    _open_tag(p)
    ea = _expected_actual(f)
    if ea:  # expected = what prod has, actual = what stage has; the words that differ are red
        left, right = _diff_words(*ea)
        for label, pieces, colour in (("Expected (prod): ", left, "2563EB"), ("Actual (stage): ", right, "0F766E")):
            p = _par(doc, after=0)
            _run(p, label, bold=True, colour=colour)
            _run(p, "“", colour=INK)
            for k, (text, changed) in enumerate(pieces):
                _run(p, (" " if k else "") + text, bold=changed, colour="B42318" if changed else INK)
            _run(p, "”", colour=INK)
    a = f.get("aem")
    if a:
        p = _par(doc, after=0)
        _run(p, "AEM topic  ", bold=True, colour=MUTED, size=8)
        _run(p, (a.get("topic") or "") + "  ·  ", size=8)
        if a.get("url"):
            _link(p, a["guid"], a["url"])
        else:
            _run(p, a.get("guid", ""), bold=True, colour="0F766E", size=8)
        # (the line ends with the GUID: no “near <element>” after it)
    _screenshots(doc, out, f, f.get("shots") or {}, img_w)
