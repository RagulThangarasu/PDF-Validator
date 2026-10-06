"""Classes used in a map's topics: every `outputclass` value in the DITA topics a map in AEM references,
with the elements that carry it, how often, and in how many topics.

  python -m pdfval classes <map path in AEM | XML-editor URL of the map> [--out reports/classes]

The map's topicrefs name their topics by GUID (GUID-…-en.dita); the topics are the .dita files of the
product folder whose root id is one of those GUIDs (sub-maps are followed). A value such as
"entry img-w60" is two classes on one element: each is counted on its own, the combination too.
"""
from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib import request as _rq
from urllib.parse import parse_qs, quote, urlsplit

from .metadata import Aem

_TAG = re.compile(r"<([A-Za-z][\w.-]*)((?:\s+[\w:.-]+\s*=\s*(?:\"[^\"]*\"|'[^']*'))*)\s*/?>")
_ATTR = re.compile(r"([\w:.-]+)\s*=\s*(?:\"([^\"]*)\"|'([^']*)')")


def map_path(arg: str) -> str:
    """The map's repository path: given as such, or taken from an XML-editor URL (its ditamap= / src=)."""
    if arg.startswith("http"):
        q = parse_qs(urlsplit(arg).query)
        arg = (q.get("ditamap") or q.get("src") or [""])[0]
    if not arg.lower().endswith(".ditamap"):
        raise ValueError("Give the map's path in AEM (…/Maps/<map>.ditamap) or its XML-editor URL")
    return arg


def _raw(aem: Aem, path: str) -> str:
    with _rq.urlopen(_rq.Request(aem.author + quote(path), headers={"Authorization": aem.auth}), timeout=60) as r:
        return r.read().decode("utf-8", "replace")


def _elements(xml: str):
    """(element, {attribute: value}) of every start tag (comments and CDATA left out)."""
    xml = re.sub(r"<!--.*?-->|<!\[CDATA\[.*?\]\]>", "", xml, flags=re.S)
    for m in _TAG.finditer(xml):
        yield m.group(1), {k: (a if a is not None else b) for k, a, b in _ATTR.findall(m.group(2))}


def _refs(xml: str) -> tuple[list[str], list[str]]:
    """Topic GUIDs and sub-map hrefs a map names, in map order."""
    topics, maps = [], []
    for el, at in _elements(xml):
        href = at.get("href", "")
        if not href or at.get("scope") == "external":
            continue
        name = href.split("#")[0].rsplit("/", 1)[-1]
        if name.lower().endswith(".ditamap") or at.get("format") == "ditamap":
            maps.append(name)
        elif name.lower().endswith((".dita", ".xml")):
            topics.append(re.sub(r"\.(dita|xml)$", "", name, flags=re.I))
    return topics, maps


def crawl(cfg: dict, map_arg: str, progress=None) -> dict:
    aem = Aem(cfg)
    mpath = map_path(map_arg)
    product = mpath.rsplit("/Maps/", 1)[0] if "/Maps/" in mpath else mpath.rsplit("/", 1)[0]
    listed = lambda pattern: sorted(h["jcr:path"] for h in aem.get(
        "/bin/querybuilder.json", path=product, type="dam:Asset", nodename=pattern,
        **{"p.limit": "-1", "p.hits": "selective", "p.properties": "jcr:path"})["hits"])
    # the map and its sub-maps
    wanted, seen_maps, todo = [], set(), [mpath]
    all_maps = {p.rsplit("/", 1)[-1]: p for p in listed("*.ditamap")}
    while todo:
        p = todo.pop(0)
        if p in seen_maps:
            continue
        seen_maps.add(p)
        topics, maps = _refs(_raw(aem, p))
        wanted += [t for t in topics if t not in wanted]
        todo += [all_maps[m] for m in maps if m in all_maps]
    # the product's topic files by their GUID (AEM Guides keeps it as jcr:content/fmUuid): the ones the map names
    hits = aem.get("/bin/querybuilder.json", path=product, type="dam:Asset", nodename="*.dita",
                   **{"p.limit": "-1", "p.hits": "selective", "p.properties": "jcr:path jcr:content/fmUuid"})["hits"]
    files = [h["jcr:path"] for h in hits]
    by_id = {(h.get("jcr:content") or {}).get("fmUuid", ""): h["jcr:path"] for h in hits}
    missing = [g for g in wanted if g not in by_id]
    todo_topics = [(g, by_id[g]) for g in wanted if g in by_id]
    done, failed = [0], []

    def read(item: tuple[str, str]) -> tuple[str, str, str]:
        guid, path = item
        xml = ""
        for attempt in range(3):  # a busy author answers 503 now and then
            try:
                xml = _raw(aem, path)
                break
            except Exception as e:
                if attempt == 2:
                    failed.append(f"{path}: {e}")
        done[0] += 1
        if progress:
            progress(done[0] / max(1, len(todo_topics)), f"Reading topics {done[0]}/{len(todo_topics)}")
        return guid, path, xml

    with ThreadPoolExecutor(4) as ex:
        used = [u for u in ex.map(read, todo_topics) if u[2]]

    count: Counter = Counter()                     # class -> uses
    on: dict[str, Counter] = defaultdict(Counter)  # class -> element -> uses
    topics_of: dict[str, set] = defaultdict(set)   # class -> topics
    combos: Counter = Counter()                    # (element, full outputclass value) -> uses
    elements: Counter = Counter()
    no_class: Counter = Counter()
    titles: dict[str, str] = {}                    # topic path -> its title
    in_topic: dict[str, Counter] = defaultdict(Counter)  # topic path -> class -> uses
    for guid, path, xml in used:
        t = re.search(r"<title[^>]*>(.*?)</title>", xml, re.S)
        titles[path] = " ".join(re.sub(r"<!--.*?-->|<[^>]+>", " ", t.group(1), flags=re.S).split()) if t else ""
        for el, at in _elements(xml):
            elements[el] += 1
            value = " ".join((at.get("outputclass") or "").split())
            if not value:
                no_class[el] += 1
                continue
            combos[(el, value)] += 1
            for c in value.split():
                count[c] += 1
                on[c][el] += 1
                topics_of[c].add(path)
                in_topic[path][c] += 1
    return {"map": mpath, "author": aem.author, "topics_in_map": len(wanted), "topics_read": len(used), "topics_missing": missing,
            "topics_failed": failed, "topic_files": len(files),
            "classes": [{"class": c, "uses": n, "topics": len(topics_of[c]),
                         "elements": ", ".join(f"{e} ({k})" for e, k in on[c].most_common()),
                         # every topic the class is used in, in map order: title, file, uses in that topic
                         "topic_list": [{"title": titles[p], "file": p.rsplit("/", 1)[-1], "uses": in_topic[p][c]}
                                        for _, p, _ in used if p in topics_of[c]]}
                        for c, n in sorted(count.items(), key=lambda x: (-x[1], x[0]))],
            # every topic of the map, in map order, with the classes it uses
            "topics": [{"title": titles[p], "file": p.rsplit("/", 1)[-1],
                        "classes": ", ".join(f"{c} ({k})" for c, k in sorted(in_topic[p].items(), key=lambda x: (-x[1], x[0])))}
                       for _, p, _ in used],
            "combinations": [{"element": e, "outputclass": v, "uses": n} for (e, v), n in sorted(combos.items(), key=lambda x: (x[0][0], -x[1]))],
            "elements": [{"element": e, "uses": n, "without_class": no_class.get(e, 0)} for e, n in sorted(elements.items(), key=lambda x: (-x[1], x[0]))]}


def write(res: dict, out_dir: str | Path) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    name = re.sub(r"\.ditamap$", "", res["map"].rsplit("/", 1)[-1], flags=re.I)
    for kind, cols in (("classes", ["class", "uses", "topics", "elements"]), ("combinations", ["element", "outputclass", "uses"]),
                       ("elements", ["element", "uses", "without_class"])):
        with open(out / f"{name}-{kind}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(res[kind])
    with open(out / f"{name}-class-topics.csv", "w", newline="", encoding="utf-8-sig") as f:  # one row per class and topic
        w = csv.writer(f)
        w.writerow(["class", "topic title", "topic file", "uses in the topic"])
        for c in res["classes"]:
            w.writerows([c["class"], t["title"], t["file"], t["uses"]] for t in c["topic_list"])
    with open(out / f"{name}-topics.csv", "w", newline="", encoding="utf-8-sig") as f:  # one row per topic
        w = csv.DictWriter(f, fieldnames=["title", "file", "classes"])
        w.writeheader()
        w.writerows(res["topics"])
    try:
        write_docx(res, out / f"{name}-classes.docx")
    except ImportError:  # python-docx not installed: the CSV lists are complete
        pass
    return out / f"{name}-classes.csv"


def write_docx(res: dict, path: str | Path) -> Path:
    """The lists as a Word document: the classes, the element + class combinations and the elements."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = max(sec.page_width, sec.page_height), min(sec.page_width, sec.page_height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)
    name = re.sub(r"\.ditamap$", "", res["map"].rsplit("/", 1)[-1], flags=re.I)
    doc.add_heading(f"Classes used in the topics of {name}", level=0)
    p = doc.add_paragraph()
    p.add_run("Map: ").bold = True
    p.add_run(res["map"])
    p = doc.add_paragraph()
    p.add_run(f"{res['topics_read']} of {res['topics_in_map']} topics read").bold = True
    p.add_run(f" · {len(res['classes'])} classes (outputclass values) · {len(res['combinations'])} element + class combinations · "
              f"{len(res['elements'])} elements" + (f" · {len(res['topics_missing'])} topics not found" if res["topics_missing"] else "")
              + (f" · {len(res['topics_failed'])} topics could not be read" if res.get("topics_failed") else ""))

    def table(title: str, note: str, heads: list[str], rows: list[list], widths: list[float]) -> None:
        doc.add_heading(title, level=1)
        if note:
            doc.add_paragraph(note).runs[0].font.color.rgb = RGBColor(0x6A, 0x72, 0x82)
        t = doc.add_table(rows=1, cols=len(heads))
        t.style = "Light Grid Accent 1"
        for cell, h in zip(t.rows[0].cells, heads):
            cell.text = ""
            cell.paragraphs[0].add_run(h).bold = True
        for r in rows:
            for cell, v in zip(t.add_row().cells, r):
                cell.text = str(v)
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)

    table(f"Classes ({len(res['classes'])})", "Every class, most used first. A value such as “entry img-w60” counts as two classes.",
          ["#", "Class", "Uses", "Topics", "On elements (uses)"],
          [[k, c["class"], c["uses"], c["topics"], c["elements"]] for k, c in enumerate(res["classes"], 1)], [0.4, 2.3, 0.7, 0.7, 5.6])
    # each class on its own: what it is on, how it is written, and every topic that uses it - as bullet points
    doc.add_heading(f"Each class in detail ({len(res['classes'])})", level=1)
    doc.add_paragraph("One block per class, in the order of the table above. The number in brackets after a topic is how many times "
                      "the class is used in that topic.").runs[0].font.color.rgb = RGBColor(0x6A, 0x72, 0x82)
    written: dict[str, list] = {}
    for c in res["combinations"]:
        for name in c["outputclass"].split():
            written.setdefault(name, []).append(c)

    def bullet(label: str, text: str = "", level: int = 1):
        p = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label).bold = True
        if text:
            p.add_run(text)
        return p

    for k, c in enumerate(res["classes"], 1):
        doc.add_heading(f"{k}. {c['class']}", level=2)
        bullet("Used: ", f"{c['uses']} time(s), in {c['topics']} topic(s)")
        bullet("On elements: ", c["elements"])
        forms = sorted(written.get(c["class"], []), key=lambda x: -x["uses"])
        bullet("Written as: ", "; ".join(f"<{x['element']} outputclass=\"{x['outputclass']}\"> ({x['uses']})" for x in forms))
        bullet(f"Topics ({c['topics']}):")
        for t in c["topic_list"]:
            p = bullet(t["title"] or "(no title)", level=2)
            r = p.add_run(f"  -  {t['file']}  ({t['uses']})")
            r.font.color.rgb = RGBColor(0x6A, 0x72, 0x82)
    table(f"Classes of each topic ({len(res['topics'])} topics)", "Every topic of the map, in map order, with the classes it uses (uses).",
          ["#", "Topic", "File", "Classes (uses)"], [[k, t["title"], t["file"], t["classes"]] for k, t in enumerate(res["topics"], 1)], [0.4, 2.4, 2.4, 4.5])
    table(f"Element + class as written ({len(res['combinations'])})", "The outputclass value exactly as it stands on each element.",
          ["Element", "outputclass", "Uses"], [[c["element"], c["outputclass"], c["uses"]] for c in res["combinations"]], [1.8, 6.6, 1.0])
    table(f"Elements ({len(res['elements'])})", "", ["Element", "Uses", "Without a class"],
          [[e["element"], e["uses"], e["without_class"]] for e in res["elements"]], [3.0, 1.5, 1.5])
    doc.save(str(path))
    return Path(path)
