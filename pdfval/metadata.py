"""Metadata check: the Document Title and Page Title of every product map in AEM against the
migration sheet (metadata/*.xlsx, sheet "Model-list").

  Document Title (dc:title)      = "{Product line} {Model name} user manual"
  Page Title     (dc:pageTitle)  = "{Product line} {Model name}"

{Product line} is the sheet's "Product line / product category" column (e.g. Monitor) and {Model name}
its "Model name" column (e.g. GW2291). The maps are every *.ditamap under <search_root>/<lang>/ in AEM
(<Brand>/<Category>/<product folder>/Maps/<map>.ditamap); a map is matched to its sheet row by the product
folder / map name against the row's model name, file version or models involved.

  python -m pdfval metadata [--sheet FILE.xlsx] [--out reports/metadata]
"""
from __future__ import annotations

import base64
import csv
import json
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html import escape
from pathlib import Path
from urllib import request as _rq
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode

import pymupdf

PROJECT = Path(__file__).resolve().parents[1]
SHEET = "Model-list"
# header text (lower case, start of the cell) -> field
COLUMNS = {"product line / product category": "category", "product line": "bu", "brand": "brand",
           "model name": "model", "models involved": "models", "file version (for metadata)2": "file",
           "file version (for metadata & output)": "version"}


# dc:description of every map (the template AEM Guides starts from: "[Product] [model]")
DESCRIPTION = "Learn how to set up {line} {model} and its settings, and optimize performance."


def default_sheet() -> Path | None:
    """The newest .xlsx in <project>/metadata."""
    files = sorted((PROJECT / "metadata").glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return next((p for p in files if not p.name.startswith("~$")), None)


def load_sheet(path: str | Path) -> list[dict]:
    """Rows of the Model-list sheet that have a model name, with the expected titles."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head_i = next(i for i, r in enumerate(rows) if any(str(c or "").strip().lower().startswith("model name") for c in r))
    cols: dict[str, int] = {}
    for j, c in enumerate(rows[head_i]):
        h = re.sub(r"\s+", " ", str(c or "")).strip().lower()
        for key, field in COLUMNS.items():  # longest header first, so "product line / ..." beats "product line"
            if h.startswith(key) and field not in cols and (field != "bu" or h == "product line"):
                if field == "version" and "2" in h:
                    continue
                cols[field] = j
                break
    out, skipped = [], 0
    for n, r in enumerate(rows[head_i + 1:], head_i + 2):
        val = lambda f: re.sub(r"\s+", " ", str(r[cols[f]])).strip() if f in cols and cols[f] < len(r) and r[cols[f]] is not None else ""
        row = {f: val(f) for f in COLUMNS.values()}
        if not row["model"]:
            skipped += bool(row["models"] or row["category"])  # planned models without a UM / model name yet
            continue
        row["row"] = n
        raw = {f: r[cols[f]] for f in ("category", "model") if f in cols and cols[f] < len(r)}
        row["sheet_notes"] = [f"sheet {f} {v!r} has extra spaces" for f, v in raw.items()
                              if isinstance(v, str) and (v != v.strip() or "  " in v)]
        row["exp_doc"] = f"{row['category']} {row['model']} user manual"
        row["exp_page"] = f"{row['category']} {row['model']}"
        row["exp_desc"] = DESCRIPTION.format(line=row["category"], model=row["model"])
        out.append(row)
    load_sheet.skipped = skipped
    return out


def _norm(s: str) -> str:
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[_-]((en|um|ug|em)[_-]?)?v\d+(\.\d+)*.*$", "", s)  # SW272_EN_V5 -> sw272 (not GV32 -> g)
    s = re.sub(r"[_\-]+", " ", s)  # Stylus_UM_EN -> "stylus um en": "_" is a word character, \b would not split there
    s = re.sub(r"\b(and|series|um|en|table)\b", " ", s)
    return re.sub(r"[^a-z0-9]", "", s)


def _tokens(s: str) -> frozenset:
    s = re.sub(r"[^a-z0-9]+", " ", s.lower().replace("&", " "))
    return frozenset(t for t in s.split() if t not in ("and", "series", "um", "en", "table"))


def models_of(cell: str) -> list[str]:
    """The sheet's "Models involved" cell as its models: "PD2706U, PD2706UA", "GW90C/GW90TC",
    "PT03,PT06; TPY24" -> each model once, in sheet order."""
    out = []
    for m in re.split(r"\s*[,;/\n]\s*|\s+(?:and|&)\s+", cell or ""):
        m = m.strip()
        if m and _norm(m) and _norm(m) not in {_norm(x) for x in out}:
            out.append(m)
    return out


def match(folder: str, map_name: str, rows: list[dict], title: str = "") -> dict | None:
    """The sheet row of a product map: exact (normalised) model name / file name first, then any one model
    of the row's "Models involved" (a manual covering several models is filed under one of them), then
    the same words in any order."""
    keys = {_norm(folder), _norm(map_name)} - {""}
    for field in ("model", "file", "models"):
        for r in rows:
            if _norm(r[field]) in keys:
                return r
    # the map's own Document Title is the row's expected title ("IFP accessory BenQ Board Pens user manual"
    # in the folder "stylus"): the map of that row, whatever its folder is called
    if title.strip():
        t = re.sub(r"\s+", " ", title).strip().lower()
        for r in rows:
            if t == (r.get("exp_doc") or "").lower():
                return r
    # a model in the row's Models involved, or a row whose model name starts with the folder's model
    # ("MA270S safety" for folder ma270s): when several rows qualify, the one whose model name the
    # map's own Document Title spells out (a safety sheet and the series manual share a model)
    hits = [r for r in rows if any(_norm(m) in keys for m in models_of(r["models"]))
            or any(k and _norm(r["model"]).startswith(k) for k in keys)]
    if hits:
        tt = _tokens(title)
        named = [r for r in hits if tt and _tokens(r["model"]) and _tokens(r["model"]) <= tt]
        return max(named, key=lambda r: len(_tokens(r["model"]))) if named else hits[0]
    toks = {_tokens(folder), _tokens(map_name)}
    for field in ("model", "file"):
        for r in rows:
            if _tokens(r[field]) in toks:
                return r
    return None


class Aem:
    def __init__(self, cfg: dict):
        self.author = (cfg.get("author") or "").rstrip("/")
        if not self.author or not cfg.get("user") or not cfg.get("password"):
            raise ValueError("Log in to AEM first (author URL, user and password)")
        self.auth = "Basic " + base64.b64encode(f"{cfg['user']}:{cfg['password']}".encode()).decode()
        self.root = cfg.get("search_root") or "/content/dam"

    def get(self, _url: str, **q) -> dict:
        url = self.author + quote(_url) + (f"?{urlencode(q)}" if q else "")
        with _rq.urlopen(_rq.Request(url, headers={"Authorization": self.auth}), timeout=30) as r:
            return json.load(r)

    def maps(self, lang: str) -> list[str]:
        root = f"{self.root.rstrip('/')}/{lang}" if lang else self.root
        hits = self.get("/bin/querybuilder.json", path=root, type="dam:Asset", nodename="*.ditamap",
                        **{"p.limit": "-1", "p.hits": "selective", "p.properties": "jcr:path"})["hits"]
        return sorted(h["jcr:path"] for h in hits)

    def languages(self) -> list[str]:
        """The language folders under the search root (en, zh-cn, ar-me, ...)."""
        try:
            kids = self.get(f"{self.root.rstrip('/')}.1.json")
        except (HTTPError, URLError, OSError, ValueError):
            return []
        return sorted(k for k, v in kids.items() if isinstance(v, dict) and re.fullmatch(r"[a-z]{2}(-[a-z]{2,4})?", k))

    def metadata(self, map_path: str) -> dict:
        try:
            return self.get(f"{map_path}/jcr:content/metadata.json")
        except HTTPError as e:
            return {"_error": f"HTTP {e.code}"}
        except (URLError, OSError, ValueError) as e:
            return {"_error": str(e)}


def _compare(actual: str, expected: str) -> str:
    if actual == expected:
        return "pass"
    if not actual:
        return "missing"
    if re.sub(r"\s+", " ", actual).strip().lower() == expected.lower():
        return "case"  # same words, different case / spacing
    if "[product]" in actual.lower() or "[model]" in actual.lower():
        return "placeholder"
    return "fail"


def _version(actual: str, expected: str) -> str:
    """dc:version against the sheet's file version, compared as numbers: V1.01 = V 1.1 = v1.1 = V1.1.0
    (format only: "case"), V1.2 ≠ V1.5. The sheet writes 1.01 as V1.1 (its file names say _V1.01)."""
    if not expected or actual == expected:
        return "pass"
    if not actual:
        return "missing"
    nums = lambda v: [int(x) for x in re.findall(r"\d+", v)] if re.fullmatch(r"\s*[vV]?\.?\s*[\d.]+\s*", v) else None
    a, e = nums(actual), nums(expected)
    if a is None or e is None:
        return "fail"
    while len(a) > 1 and a[-1] == 0:
        a.pop()
    while len(e) > 1 and e[-1] == 0:
        e.pop()
    return "case" if a == e else "fail"


def _near(a: str, b: str) -> bool:
    """A spelling slip (1-2 letters) rather than a different name."""
    from difflib import SequenceMatcher
    a, b = a.lower().strip(), b.lower().strip()
    return a != b and len(a) > 6 and SequenceMatcher(None, a, b).ratio() >= 0.93


def validate(cfg: dict, sheet: str | Path | None = None, lang: str = "en", progress=None) -> dict:
    """Pull dc:title / dc:pageTitle of every product map and check them against the sheet."""
    sheet = Path(sheet) if sheet else default_sheet()
    if not sheet or not sheet.is_file():
        raise ValueError("No metadata sheet (.xlsx) found in the metadata folder")
    rows = load_sheet(sheet)
    aem = Aem(cfg)
    if progress:
        progress(0.02, "Listing product maps in AEM")
    paths = aem.maps(lang)
    done = [0]

    def one(p: str, lang: str = lang, rows: list[dict] = rows, force: dict | None = None) -> dict:
        md = aem.metadata(p)
        done[0] += 1
        if progress:
            progress(0.05 + 0.9 * min(1, done[0] / max(1, len(paths))), f"Reading map metadata {done[0]}/{len(paths)}")
        parts = p.split("/")
        mi = parts.index("Maps") if "Maps" in parts else len(parts) - 1
        folder = parts[mi - 1]
        rel = "/".join(parts[parts.index(lang) + 1:mi] if lang in parts else parts[:mi])
        row = force or match(folder, re.sub(r"\.ditamap$", "", parts[-1], flags=re.I), rows, str(md.get("dc:title") or ""))
        item = {"map": parts[-1], "path": p, "product": rel, "folder": folder,
                "doc": str(md.get("dc:title") or ""), "page": str(md.get("dc:pageTitle") or ""),
                "desc": str(md.get("dc:description") or ""), "version": str(md.get("dc:version") or ""),
                "product_line": str(md.get("dc:productLine") or ""), "error": md.get("_error", ""),
                "url": f"{aem.author}/assetdetails.html{quote(p)}"}
        if row:
            item.update(row=row["row"], model=row["model"], category=row["category"], brand=row["brand"], file=row["file"],
                        exp_doc=row["exp_doc"], exp_page=row["exp_page"],
                        doc_status=_compare(item["doc"], row["exp_doc"]), page_status=_compare(item["page"], row["exp_page"]),
                        line_ok=_norm(item["product_line"]) == _norm(row["category"]),
                        exp_desc=row["exp_desc"], desc_status=_compare(item["desc"], row["exp_desc"]),
                        exp_version=row["version"], version_status=_version(item["version"], row["version"]),
                        notes=list(row["sheet_notes"]))
            if (item["doc_status"] == "fail" and _near(item["doc"], row["exp_doc"])) or \
                    (item["page_status"] == "fail" and _near(item["page"], row["exp_page"])):
                item["notes"].append("1-2 letters differ: check the spelling in AEM and in the sheet")
            if "timing" in (folder + parts[-1]).lower() and "timing" not in row["model"].lower():
                item["notes"].append(f"timing-table map, but the sheet's model name is {row['model']!r} (no 'timing')")
            sheet_models = models_of(row["models"])
            keys = {_norm(folder), _norm(re.sub(r"\.ditamap$", "", parts[-1], flags=re.I))}
            if not keys & {_norm(row["model"]), _norm(row["file"])}:
                by = next((x for x in sheet_models if _norm(x) in keys), None)
                if by:
                    item["notes"].append(f"matched by {by!r} in the sheet's Models involved ({', '.join(sheet_models)})")
            item["exp_models"] = ", ".join(sheet_models)
            dup = [m for m in re.split(r"\s*[,;/]\s*", row["models"] or "") if m.strip()]
            if len(dup) > len(sheet_models):
                item["notes"].append(f"sheet 'Models involved' lists a model twice: {row['models']!r}")
            if re.search(r"test", rel, re.I):
                item["notes"].append("test / copy folder")
            checked = (item["doc_status"], item["page_status"], item["desc_status"])
            item["status"] = ("error" if item["error"] else
                              "pass" if all(x == "pass" for x in checked) else
                              "fail" if any(x in ("fail", "missing", "placeholder") for x in checked) else "warn")
        else:
            item["status"] = "unmatched"
        return item

    with ThreadPoolExecutor(8) as ex:
        items = list(ex.map(one, paths))
    # a second map on a sheet row that already has its own map, while a row with a near-identical file name
    # has none: folder "g90" beside "G90-Series" (row "G90 Series"), and the row "GW90C Series" (file
    # G90C-EM-V1) without a map - the second map is that row's map, filed under a shortened name
    def reassign() -> None:
        taken = {i.get("row") for i in items}
        free = [r for r in rows if r["row"] not in taken]
        groups: dict[int, list[int]] = {}
        for k, i in enumerate(items):
            if i.get("row"):
                groups.setdefault(i["row"], []).append(k)
        for ks in groups.values():
            if len(ks) < 2 or not any(items[k]["doc_status"] == "pass" for k in ks):
                continue  # one map, or none of them is clearly the row's own map
            for k in ks:
                i = items[k]
                if i["doc_status"] == "pass" or re.search(r"test", i["product"], re.I):
                    continue
                key = _norm(i["folder"])
                hit = [r for r in free if key and any(_norm(r[f]).startswith(key) and len(_norm(r[f])) - len(key) <= 2 for f in ("file", "model"))]
                if len(hit) == 1:
                    old = i["model"]
                    items[k] = one(i["path"], lang, rows, force=hit[0])
                    items[k]["notes"].append(f"folder {i['folder']!r} taken as this row's map (file {hit[0]['file']!r}); "
                                             f"the row {old!r} has its own map")
                    free.remove(hit[0])

    reassign()
    by_row: dict[int, list[dict]] = {}
    for i in items:
        if i.get("row"):
            by_row.setdefault(i["row"], []).append(i)
    for same in by_row.values():
        for i in same if len(same) > 1 else ():
            i["notes"].append("same sheet row as " + ", ".join(o["product"] + "/" + o["map"] for o in same if o is not i))
    seen = {i.get("row") for i in items}
    missing = [r for r in rows if r["row"] not in seen]
    # a model with no map under /<lang>/ may have its map under another language folder (a manual whose
    # source is Chinese: i800_i800ST_UM_ZH-CN under /zh-cn/): the map is in AEM - validated there, and noted
    if missing:
        for other in [x for x in aem.languages() if x != lang]:
            try:
                other_paths = aem.maps(other)
            except (HTTPError, URLError, OSError, ValueError):
                continue
            for p in other_paths:
                item = one(p, other, missing)
                if item["status"] == "unmatched":
                    continue
                item["notes"].append(f"the map is under /{other}/, not /{lang}/")
                item["lang"] = other
                items.append(item)
        seen = {i.get("row") for i in items}
        missing = [r for r in rows if r["row"] not in seen]
    count = lambda s: sum(i["status"] == s for i in items)
    return {"created": datetime.now().isoformat(timespec="seconds"), "sheet": sheet.name, "author": aem.author,
            "lang": lang, "maps": items, "not_in_aem": missing,
            "summary": {"maps": len(items), "pass": count("pass"), "warn": count("warn"), "fail": count("fail"),
                        "unmatched": count("unmatched"), "error": count("error"), "not_in_aem": len(missing),
                        "desc_bad": sum(i.get("desc_status") not in (None, "pass") for i in items),
                        "desc_placeholder": sum(i.get("desc_status") == "placeholder" for i in items),
                        "line_bad": sum(i.get("line_ok") is False for i in items),
                        "version_bad": sum(i.get("version_status") in ("fail", "missing") for i in items),
                        "notes": sum(bool(i.get("notes")) for i in items),
                        "rows": len(rows), "rows_skipped": getattr(load_sheet, "skipped", 0)}}


# ---------------------------------------------------------------- reports

STATUS = {"pass": ("Pass", "#16a34a"), "case": ("Case / spacing", "#b45309"), "fail": ("Mismatch", "#d92d20"),
          "missing": ("Empty in AEM", "#d92d20"), "placeholder": ("Mismatch", "#d92d20"),
          "not_in_aem": ("No map in AEM", "#64748b"), "error": ("Not readable", "#d92d20")}
ORDER = {"fail": 0, "missing": 1, "error": 2, "case": 3, "not_in_aem": 4, "pass": 5}
CSS = """
* { font-family: sans-serif; font-size: 8.5px; color: #1d2330; }
h1 { font-size: 18px; margin: 0 0 2px 0; } h2 { font-size: 12px; margin: 14px 0 4px 0; }
table { border-collapse: collapse; width: 100%; }
th, td { border-bottom: 0.5px solid #d0d4da; padding: 3px 4px; text-align: left; vertical-align: top; }
th { background-color: #f1f3f6; font-weight: bold; }
.muted { color: #6a7282; } .n { text-align: right; }
"""


def _field(actual: str, expected: str) -> str:
    """pass / case (only capitals or spaces differ) / missing (empty in AEM) / fail."""
    if actual == expected:
        return "pass"
    if not actual.strip():
        return "missing"
    if re.sub(r"\s+", " ", actual).strip().lower() == expected.lower():
        return "case"
    return "fail"


def _worst(*ss: str) -> str:
    return min(ss, key=lambda x: ORDER[x])


def sheet_rows(res: dict) -> list[dict]:
    """The Excel is the base: one entry per sheet row (sheet order), with the AEM map(s) found for it and
    the Document Title / Page Title result of each."""
    by_row: dict[int, dict] = {}
    for i in res["maps"]:
        if not i.get("row"):
            continue
        r = by_row.setdefault(i["row"], {"row": i["row"], "brand": i.get("brand", ""), "category": i["category"],
                                         "model": i["model"], "file": i.get("file", ""), "exp_doc": i["exp_doc"],
                                         "exp_page": i["exp_page"], "exp_desc": i.get("exp_desc", ""),
                                         "exp_models": i.get("exp_models", ""), "maps": []})
        d, p, ds = _field(i["doc"], i["exp_doc"]), _field(i["page"], i["exp_page"]), _field(i.get("desc", ""), i.get("exp_desc", ""))
        r["maps"].append({**{k: i[k] for k in ("product", "map", "path", "url", "doc", "page")}, "desc": i.get("desc", ""),
                          "doc_status": d, "page_status": p, "desc_status": ds, "notes": i.get("notes", []),
                          "status": "error" if i.get("error") else _worst(d, p, ds),
                          "error": i.get("error", "")})
    for r in res["not_in_aem"]:
        by_row[r["row"]] = {**{k: r.get(k, "") for k in ("row", "brand", "category", "model", "file", "exp_doc", "exp_page", "exp_desc")},
                            "exp_models": ", ".join(models_of(r.get("models", ""))), "maps": []}
    out = sorted(by_row.values(), key=lambda r: r["row"])
    for r in out:
        # a row with several maps (copies, test folders) is as bad as its worst map: every copy is published,
        # so a mismatched copy is a mismatch even when another copy is right (TEY41 in TWY31/ and tey41/)
        r["status"] = _worst(*(m["status"] for m in r["maps"])) if r["maps"] else "not_in_aem"
    return out


def _summary(rows: list[dict], res: dict) -> dict:
    c = lambda s: sum(r["status"] == s for r in rows)
    return {"rows": len(rows), "pass": c("pass"), "fail": c("fail"), "missing": c("missing"), "case": c("case"),
            "not_in_aem": c("not_in_aem"), "error": c("error"),
            "doc_ok": sum(any(m["doc_status"] == "pass" for m in r["maps"]) for r in rows),
            "page_ok": sum(any(m["page_status"] == "pass" for m in r["maps"]) for r in rows),
            "desc_ok": sum(any(m["desc_status"] == "pass" for m in r["maps"]) for r in rows),
            "extra_maps": sum(i["status"] == "unmatched" for i in res["maps"]), "skipped": res["summary"].get("rows_skipped", 0)}


def _badge(s: str) -> str:
    label, color = STATUS[s]
    return f'<b style="color:{color}">{escape(label)}</b>'


def _cell(actual: str, status: str) -> str:
    if status == "pass":
        return escape(actual)
    return f'<span style="color:{STATUS[status][1]}">{escape(actual) or "<i>(empty)</i>"}</span>'


def _wrap(t: str) -> str:
    """Escaped text that can wrap in a narrow cell: a space after every / and , (a long run without spaces
    cannot be broken and never fits the table)."""
    return escape(re.sub(r"([/,])(?=\S)", r"\1 ", t))


def _models_line(r: dict) -> str:
    """The sheet's Models involved under the model name (the models the manual covers)."""
    return f'<br/><span class="muted">Models: {_wrap(r["exp_models"])}</span>' if r.get("exp_models") else ""


def build_pdf(res: dict, out: str | Path) -> Path:
    rows = sheet_rows(res)
    sm = _summary(rows, res)
    html = [f"""<h1>AEM map titles vs Excel</h1>
<p class="muted">The Excel is the base. For every model in the sheet the product map in AEM must have<br/>
<b>Document Title</b> (dc:title) = "{{Product line}} {{Model name}} user manual", <b>Page Title</b> (dc:pageTitle) = "{{Product line}} {{Model name}}"<br/>
and <b>Meta Description</b> (dc:description) = "{escape(DESCRIPTION.format(line='{Product line}', model='{Model name}'))}"<br/>
{{Product line}} = column "Product line / product category", {{Model name}} = column "Model name" · the match is exact (capitals and spaces count).<br/>
Excel {escape(res['sheet'])} (sheet Model-list) · AEM {escape(res['author'])} /{escape(res['lang'])}/ · {escape(res['created'].replace('T', ' '))}</p>
<table><tr><th>Excel rows</th><th>Pass</th><th>Mismatch</th><th>Empty in AEM</th><th>Case / spacing only</th><th>No map in AEM</th>
<th>Document Title correct</th><th>Page Title correct</th><th>Meta Description correct</th></tr>
<tr><td>{sm['rows']}</td><td style="color:#16a34a"><b>{sm['pass']}</b></td><td style="color:#d92d20"><b>{sm['fail']}</b></td>
<td style="color:#d92d20">{sm['missing']}</td><td style="color:#b45309">{sm['case']}</td><td>{sm['not_in_aem']}</td>
<td>{sm['doc_ok']} / {sm['rows']}</td><td>{sm['page_ok']} / {sm['rows']}</td><td>{sm['desc_ok']} / {sm['rows']}</td></tr></table>
<p class="muted">{sm['skipped']} sheet rows have no model name yet (planned models, "naming TBD") and are not in scope.</p>"""]

    def table(title: str, sel: list[dict]) -> None:
        if not sel:
            return
        html.append(f"""<h2>{escape(title)} ({len(sel)})</h2><table><tr><th>Row</th><th>Product line</th><th>Model name</th>
<th>Result</th><th>Expected (Excel)</th><th>In AEM</th><th>AEM map</th></tr>""")
        for r in sel:
            exp = (f'<span class="muted">Doc:</span> {escape(r["exp_doc"])}<br/><span class="muted">Page:</span> {escape(r["exp_page"])}'
                   f'<br/><span class="muted">Desc:</span> {escape(r.get("exp_desc", ""))}')
            if not r["maps"]:
                html.append(f"<tr><td class='n'>{r['row']}</td><td>{escape(r['category'])}</td><td><b>{escape(r['model'])}</b></td>"
                            f"<td>{_badge('not_in_aem')}</td><td>{exp}</td><td class='muted'>—</td><td class='muted'>file {escape(r['file'])}</td></tr>")
                continue
            for k, m in enumerate(r["maps"]):
                head = (f"<td class='n'>{r['row']}</td><td>{escape(r['category'])}</td><td><b>{escape(r['model'])}</b></td>"
                        if k == 0 else "<td></td><td></td><td></td>")
                got = (f'<span class="muted">Doc:</span> {_cell(m["doc"], m["doc_status"])}<br/>'
                       f'<span class="muted">Page:</span> {_cell(m["page"], m["page_status"])}'
                       + f'<br/><span class="muted">Desc:</span> {_cell(m["desc"], m["desc_status"])}'
                       + "".join(f'<br/><span class="muted">· {_wrap(n)}</span>' for n in m.get("notes", [])))
                html.append(f"<tr>{head}<td>{_badge(m['status'])}</td><td>{exp if k == 0 else ''}</td><td>{got}</td>"
                            f"<td>{_wrap(m['product'])}<br/><span class='muted'>{escape(m['map'])}</span></td></tr>")
        html.append("</table>")

    bad = sorted([r for r in rows if r["status"] in ("fail", "missing", "error", "case")], key=lambda r: (ORDER[r["status"]], r["row"]))
    table("Rows to fix in AEM", bad)
    table("Models in the Excel with no map in AEM", [r for r in rows if r["status"] == "not_in_aem"])
    table("Rows that pass", [r for r in rows if r["status"] == "pass"])
    extra = [i for i in res["maps"] if i["status"] == "unmatched"]
    if extra:
        html.append(f"""<h2>AEM maps not in the Excel ({len(extra)})</h2><p class="muted">Not validated: no model in the sheet for them.</p>
<table><tr><th>AEM map</th><th>Document Title</th><th>Page Title</th></tr>""")
        html += [f"<tr><td>{_wrap(i['product'])}<br/><span class='muted'>{escape(i['map'])}</span></td>"
                 f"<td>{escape(i['doc']) or '<i>(empty)</i>'}</td><td>{escape(i['page']) or '<i>(empty)</i>'}</td></tr>" for i in extra]
        html.append("</table>")

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    story = pymupdf.Story("".join(html), user_css=CSS)
    writer = pymupdf.DocumentWriter(str(out))
    page = pymupdf.paper_rect("a4-l")
    where = page + (28, 28, -28, -36)
    more, n = 1, 0
    while more and n < 400:  # a row that fits no page would add pages forever
        dev = writer.begin_page(page)
        more, filled = story.place(where)
        story.draw(dev)
        writer.end_page()
        n += 1
    writer.close()
    doc = pymupdf.open(out)
    for k, pg in enumerate(doc):
        pg.insert_text((page.width - 90, page.height - 16), f"Page {k + 1} / {len(doc)}", fontsize=7, color=(0.42, 0.45, 0.51))
    doc.saveIncr()
    return out


def write_csv(res: dict, out: str | Path) -> Path:
    out = Path(out)
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Excel row", "Product line", "Model name", "Result", "Expected Document Title", "Document Title in AEM",
                    "Document Title result", "Expected Page Title", "Page Title in AEM", "Page Title result",
                    "Expected Meta Description", "Meta Description in AEM", "Meta Description result", "Notes", "AEM map", "AEM link"])
        for r in sheet_rows(res):
            for m in r["maps"] or [None]:
                w.writerow([r["row"], r["category"], r["model"], STATUS[m["status"] if m else "not_in_aem"][0],
                            r["exp_doc"], m["doc"] if m else "", STATUS[m["doc_status"]][0] if m else "",
                            r["exp_page"], m["page"] if m else "", STATUS[m["page_status"]][0] if m else "",
                            r.get("exp_desc", ""), m["desc"] if m else "", STATUS[m["desc_status"]][0] if m else "",
                            "; ".join(m.get("notes", [])) if m else "", m["path"] if m else "", m["url"] if m else ""])
    return out


def run(cfg: dict, out_dir: str | Path, sheet: str | Path | None = None, lang: str = "en", progress=None) -> dict:
    res = validate(cfg, sheet, lang, progress)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "metadata.json").write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    build_pdf(res, out_dir / "metadata-report.pdf")
    write_csv(res, out_dir / "metadata-report.csv")
    rows = sheet_rows(res)
    res["excel"], res["excel_rows"] = _summary(rows, res), rows
    (out_dir / "metadata.json").write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    return res
